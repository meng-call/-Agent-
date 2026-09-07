"""
合约风控审查 Agent 系统 — DOCX 文档解析器.

基于 python-docx 提取 Word 文档中的文本内容。
"""

import logging
from pathlib import Path

from app.parsers.base import BaseParser

logger = logging.getLogger(__name__)


class DocxParser(BaseParser):
    """DOCX 文档解析器，提取段落纯文本."""

    async def parse(self, file_path: Path) -> str:
        """解析 DOCX 文件，返回提取的纯文本.

        Args:
            file_path: DOCX 文件路径.

        Returns:
            提取的纯文本，段落间以双换行分隔.

        Raises:
            FileNotFoundError: 文件不存在.
            ValueError: 文件不是 DOCX 格式.
        """
        if not file_path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")

        suffix = file_path.suffix.lower()
        if suffix not in (".docx",):
            raise ValueError(f"不支持的文件格式: {suffix}，需要 .docx")

        try:
            from docx import Document
        except ImportError:
            raise ImportError(
                "python-docx 未安装，请执行: pip install python-docx>=1.1.0"
            )

        try:
            doc = Document(str(file_path))

            text_parts: list[str] = []
            for para in doc.paragraphs:
                text = para.text
                if text and text.strip():
                    text_parts.append(text.strip())

            # 也提取表格中的文本
            for table in doc.tables:
                for row in table.rows:
                    row_texts: list[str] = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_texts.append(cell.text.strip())
                    if row_texts:
                        text_parts.append(" | ".join(row_texts))

            extracted = "\n\n".join(text_parts)
            logger.info(
                f"DocxParser: 提取完成, {len(extracted)} 字符, "
                f"{len(doc.paragraphs)} 段落, {len(doc.tables)} 表格"
            )
            return extracted

        except Exception as e:
            logger.error(f"DocxParser: 解析失败 → {file_path.name}: {e}")
            raise RuntimeError(f"DOCX 解析失败: {e}") from e
